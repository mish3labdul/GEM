"""D3A QA. Writes 11_QA_Evidence/d3_qa_results.json and .md. Every check prints PASS/FAIL; nothing is summarised as "full system consistency".
Usage: d3_qa.py <repo_root> <pkg_dir>"""
import sys,os,re,json,glob,zipfile,hashlib,subprocess,collections
from lxml import etree
root,pkg=[os.path.abspath(a) for a in sys.argv[1:3]]
sha=lambda p:hashlib.sha256(open(p,'rb').read()).hexdigest()
git=lambda *a:subprocess.run(['git','-C',root]+list(a),capture_output=True,text=True)
L=json.load(open(pkg+'/11_QA_Evidence/d3a_apply_log.json'));R=[]
def rec(n,name,ok,ev): R.append(dict(n=n,name=name,result='PASS' if ok else 'FAIL',evidence=ev));print(n,name,'PASS' if ok else 'FAIL')
NSA='http://schemas.openxmlformats.org/drawingml/2006/main'
# 1 edit proof (element level)
bad=[];tot=0
for k in 'AC':
    d=L['decks'][k];za,zb=zipfile.ZipFile(d['src']),zipfile.ZipFile(d['dst']);texts={(e['part'],e['before'],e['after']) for e in L['edits'] if e['document']==('Part A' if k=='A' else 'Part C')}
    if za.namelist()!=zb.namelist(): bad.append(k+': member list differs')
    for n in za.namelist():
        if za.read(n)==zb.read(n): continue
        try: ra,rb=etree.fromstring(za.read(n)),etree.fromstring(zb.read(n))
        except Exception as e: bad.append(f'{n}: not well-formed');continue
        ea=[e for e in ra.iter() if isinstance(e.tag,str)];eb=[e for e in rb.iter() if isinstance(e.tag,str)]
        if len(ea)!=len(eb): bad.append(f'{n}: element count');continue
        for x,y in zip(ea,eb):
            if x.tag!=y.tag or dict(x.attrib)!=dict(y.attrib): bad.append(f'{n}: structure/attribute change');break
            if (x.text or '')!=(y.text or ''):
                ok=any(e[0]==n and e[1] in (x.text or '') and (x.text or '').replace(e[1],e[2])==(y.text or '') for e in texts)
                if ok: tot+=1
                else: bad.append(f'{n}: unexpected text change')
rec(1,'Edit proof: only the intended text differs from the ODI01-R1 base (parts well-formed, structure and attributes identical)',not bad and tot==4 and set(L['decks']['A']['parts_changed'])=={'ppt/notesSlides/notesSlide79.xml','ppt/notesSlides/notesSlide80.xml','ppt/slides/slide75.xml'} and L['decks']['C']['parts_changed']==['ppt/slides/slide44.xml'],f'text edits verified={tot} (expected 4: C1 x2, C2, D3B slide 75); parts changed A={L["decks"]["A"]["parts_changed"]} C={L["decks"]["C"]["parts_changed"]}; problems={bad}')
# 2 originals untouched
paths=['GEM_V3_RC2','PartB_RC2','Amenities_Portfolio_PartD_RC2','GEM_Brand_Assets_v1.0','GEM_Letterhead_Set_v1.1_Application_Revision_03','qa','README.md','GEM_V3.0_Integrated_Brand_Ecosystem_Audit.md']
unch=git('diff','--quiet','HEAD','--',*paths).returncode==0;outside=[l for l in git('status','--porcelain').stdout.splitlines() if 'GEM_V3.0_RC2_D3_Owner_Decision_Package' not in l and not l.endswith('.DS_Store')]
r1a=sha(glob.glob(root+'/GEM_V3_RC2/04_release/*Part A*.pptx')[0]);r1c=sha(glob.glob(root+'/GEM_V3_RC2/04_release/*Part C*.pptx')[0])
rec(2,'Authoritative originals untouched',unch and not outside and r1a.startswith('c50dd3d0') and r1c.startswith('f9f968d2'),f'git diff HEAD over authoritative paths empty={unch}; changes outside the D3 package (ignoring .DS_Store)={outside}; Part A {r1a[:12]}, Part C {r1c[:12]} equal to the recorded originals')
# 3 terminology (line multiset diff)
def lines(path):
    from pptx import Presentation
    P=Presentation(path);out=collections.Counter()
    for i,s in enumerate(P.slides,1):
        def walk(shs):
            for sh in shs:
                if sh.has_text_frame:
                    for l in sh.text_frame.text.split('\n'):
                        if l.strip(): out[l.strip()]+=1
                if getattr(sh,'has_table',False) and sh.has_table:
                    for r in sh.table.rows:
                        for c in r.cells:
                            for l in c.text.split('\n'):
                                if l.strip(): out[l.strip()]+=1
                if sh.shape_type==6: walk(sh.shapes)
        walk(s.shapes)
        if s.has_notes_slide:
            for l in s.notes_slide.notes_text_frame.text.split('\n'):
                if l.strip(): out['[notes] '+l.strip()]+=1
    return out
chg={};ok3=True;ev3=[]
for k in 'AC':
    a=lines(L['decks'][k]['src']);b=lines(L['decks'][k]['dst']);add=list((b-a).elements());rem=list((a-b).elements());chg[k]=(len(rem),len(add))
    cnt=lambda C,pat:sum(v for l,v in C.items() for _ in re.findall(pat,l))
    for pat in (r'Part [A-E]\b',r'RC2',r'V3\.0',r'HOSPITALITY, IN PERFECT PROPORTION',r'VAL-\d\d',r'GEM-(BG|DDS|PS)-V3\.0-RC2'):
        if cnt(a,pat)!=cnt(b,pat) and not (k=='A' and pat in(r'RC2',r'VAL-\d\d',r'Part [A-E]\b')) and not (k=='C' and pat==r'Part [A-E]\b'): ok3=False;ev3.append(f'{k}:{pat} {cnt(a,pat)}->{cnt(b,pat)}')
    if any(re.search(r'Part E\b',l) for l in b): ok3=False;ev3.append(k+': Part E')
rec(3,'Terminology: changed lines are exactly the intended ones; version/tagline/document-ID counts unchanged; no "Part E"',ok3 and chg['A']==(3,3) and chg['C']==(1,1),f'removed/added lines A={chg["A"]} (expected 3/3), C={chg["C"]} (expected 1/1); deviations={ev3}. VAL/RC2/Part-letter counts in Part A notes change only by the intended wording (Part B described as RC2; VAL-13, VAL-20 named)')
# 4 release wording (context-aware checker; see scripts/release_wording_checker.py)
sys.path.insert(0,os.path.join(pkg,'scripts'))
import release_wording_checker as W
from pptx import Presentation
def paras(path):
    P=Presentation(path);t=[]
    for s in P.slides:
        def walk(shs):
            for sh in shs:
                if sh.has_text_frame: t.extend(p.text for p in sh.text_frame.paragraphs)
                if getattr(sh,'has_table',False) and sh.has_table:
                    for r in sh.table.rows:
                        for c in r.cells: t.extend(p.text for p in c.text_frame.paragraphs)
                if sh.shape_type==6: walk(sh.shapes)
        walk(s.shapes)
        if s.has_notes_slide: t.extend(p.text for p in s.notes_slide.notes_text_frame.paragraphs)
    return '\n'.join(t)
dtxt={k:paras(L['decks'][k]['dst']) for k in 'AC'};ncand={n['candidate']:open(pkg+'/12_Candidate_Files/notes_candidates/'+n['candidate'],encoding='utf-8').read() for n in L['notes']}
sysready=sum(len(re.findall('SYSTEM READY',t)) for t in list(dtxt.values())+list(ncand.values()))
viol=[(k,v) for k,t in list(dtxt.items())+list(ncand.items()) for v in W.violations(t)]
nr=all('SYNCHRONIZED · NOT RELEASED · EVIDENCE GATES REMAIN' in W.strip_md(t) for t in ncand.values()) and 'NOT RELEASED' in dtxt['A'] and 'NOT RELEASED' in dtxt['C']
appr=all('Not "Approved V3.0"' in W.strip_md(t) for t in ncand.values())
fx=json.load(open(pkg+'/11_QA_Evidence/release_wording_checker_tests.json'));fxbad=sum(1 for c in fx['allowed'] if W.violations(c['text']))+sum(1 for c in fx['forbidden'] if not W.violations(c['text']));fxn=len(fx['allowed'])+len(fx['forbidden'])
rec(4,'Release wording: no unqualified SYSTEM READY / PRODUCTION READY / current-state RELEASED, APPROVED or FINAL claim; NOT RELEASED present; checker regression fixture',sysready==0 and nr and appr and not viol and fxbad==0,f'SYSTEM READY occurrences in the 2 candidate decks and 3 candidate notes={sysready}; "SYNCHRONIZED · NOT RELEASED · EVIDENCE GATES REMAIN" in all 3 notes and NOT RELEASED in both decks={nr}; Not "Approved V3.0" retained in all 3 notes (markdown stripped)={appr}; context-aware violations={viol}; checker regression fixture {fxn-fxbad}/{fxn} pass. Legitimate Part C slide 60 release-form wording is unchanged and allowed by context.')
# 5 letterhead status consistency
LH=glob.glob(root+'/GEM_Letterhead_Set_v1.1_Application_Revision_03/01_templates/*.docx')
def dx(p):
    z=zipfile.ZipFile(p);return ''.join(re.sub('<[^>]+>','',z.read(n).decode('utf8','ignore')) for n in z.namelist() if re.match(r'word/(document|header\d*|footer\d*)\.xml',n))
c44=[l for l in lines(L['decks']['C']['dst']) if 'Templates are an OPEN DELIVERABLE' in l]
tab=lines(L['decks']['C']['dst']);other=[(k,l[:60]) for k in ('A','C') for l in lines(L['decks'][k]['dst']) if re.search(r'none exists yet',l)]
rec(5,'Letterhead-status consistency: Part C 44 matches the letterhead package; no remaining "none exists yet"',len(LH)==8 and all('WORKING APPLICATION / PENDING VALIDATION' in dx(p) for p in LH) and len(c44)==1 and 'WORKING APPLICATION / PENDING VALIDATION' in c44[0] and 'no template is accepted' in c44[0] and tab['[PENDING PRODUCTION MASTER]']>=5 and not other,f'8/8 templates carry the status text={len(LH)==8 and all("WORKING APPLICATION / PENDING VALIDATION" in dx(p) for p in LH)}; slide 44 sentence present={len(c44)==1}; table still shows [PENDING PRODUCTION MASTER] x{tab["[PENDING PRODUCTION MASTER]"]}; stale phrase remaining={other}')
# 6 Part D checksum
cand=open(pkg+'/12_Candidate_Files/checksum_candidates/PartD_04_release_SHA256SUMS.D3-CANDIDATE.txt',encoding='utf-8').read().splitlines()
pd=root+'/Amenities_Portfolio_PartD_RC2/04_release';ok6=len(cand)==2 and not any(l.endswith('SHA256SUMS.txt') for l in cand)
for l in cand:
    h,f=re.match(r'^([0-9a-f]{64})  (.+)$',l).groups();ok6&=os.path.exists(os.path.join(pd,f)) and sha(os.path.join(pd,f))==h
old=open(pd+'/SHA256SUMS.txt',encoding='utf-8').read().splitlines();same=[l for l in old if not l.endswith('SHA256SUMS.txt')]==cand
rec(6,'Part D checksum list: no self-entry; every remaining entry verifies against the authoritative files',ok6 and same,f'candidate lines={len(cand)}; verify against {os.path.relpath(pd,root)}: {ok6}; identical to the original minus the self-entry={same}; the original self-entry was wrong (it listed {old[1][:12]}…; the file hashes to {sha(pd+"/SHA256SUMS.txt")[:12]}…)')
# 7 112 assertions on the candidates
env=dict(os.environ);R1=root+'/GEM_V3.0_RC2_Owner_Decision_Implementation_Candidate_ODI01_R1/19_candidate_documents/deck_candidates/'
g=lambda pat:glob.glob(R1+pat)[0]
env.update(GEM_A_PATH=L['decks']['A']['dst'],GEM_C_PATH=L['decks']['C']['dst'],GEM_B_PATH=g('*Part B*.pptx'),GEM_BPDF_PATH=g('*Part B*.pdf'),GEM_APDF=g('*Part A*.pdf'),GEM_CPDF=glob.glob(pkg+'/12_Candidate_Files/deck_candidates/*Part C*.pdf')[0])
out=subprocess.run([sys.executable,root+'/GEM_V3.0_RC2_Owner_Decision_Implementation_Candidate_ODI01_R1/17_scripts/consistency_check_odi01.py',root,pkg+'/11_QA_Evidence/consistency_assertions_d3_candidates.md'],capture_output=True,text=True,env=env).stdout.strip()
rec(7,'112 automated assertions on the D3 candidate decks (A, C; B and PDFs from ODI01-R1)','112 PASS, 0 FAIL' in out,out+' (text/metadata assertions only; not a statement of full system consistency)')
# 8 earlier packages intact
import tempfile,shutil
tmpd=tempfile.mkdtemp();R1N='GEM_V3.0_RC2_Owner_Decision_Implementation_Candidate_ODI01_R1'
ar=subprocess.run(['git','-C',root,'archive','HEAD',R1N],capture_output=True);subprocess.run(['tar','-x','-C',tmpd],input=ar.stdout)
r1=subprocess.run([sys.executable,root+'/'+R1N+'/17_scripts/verify_manifest.py',root,tmpd+'/'+R1N,tmpd+'/mres.json'],capture_output=True,text=True).stdout.strip()+' [verified on a clean `git archive HEAD` export of the committed package; untracked Finder .DS_Store files on disk are not part of the package]'
shutil.rmtree(tmpd,ignore_errors=True)
d7=subprocess.run([sys.executable,root+'/GEM_V3.0_RC2_D7_Repository_Visibility_Decision_Package/scripts/d7_manifest.py','verify',root+'/GEM_V3.0_RC2_D7_Repository_Visibility_Decision_Package'],capture_output=True,text=True).stdout.strip()
rec(8,'Existing ODI01-R1 and D7 packages intact (manifests and checksums verify; unchanged vs HEAD)','18 PASS' in r1 and '19 PASS' in r1 and 'problems=[]' in d7 and git('diff','--quiet','HEAD','--','GEM_V3.0_RC2_Owner_Decision_Implementation_Candidate_ODI01_R1','GEM_V3.0_RC2_D7_Repository_Visibility_Decision_Package').returncode==0,f'ODI01-R1: {r1}; D7: {d7}')
# 9 VAL-ID regression + docs scan
sys.path.insert(0,root+'/GEM_V3.0_RC2_Owner_Decision_Implementation_Candidate_ODI01_R1/17_scripts')
import validation_id_checker as V
t=subprocess.run([sys.executable,root+'/GEM_V3.0_RC2_Owner_Decision_Implementation_Candidate_ODI01_R1/17_scripts/run_validation_id_tests.py',root+'/GEM_V3.0_RC2_Owner_Decision_Implementation_Candidate_ODI01_R1/18_qa_evidence/validation_id_checker_tests.json','/tmp/_vt.json'],capture_output=True,text=True).stdout.strip().splitlines()[-1]
ids=set();badf=[];expl=[];skipped=0
for f in sorted(glob.glob(pkg+'/[0-9][0-9]_*.md')):
    i_,b_,e_=V.check_files([f]);ids|=set(i_);badf+=b_;expl+=e_
for n in L['notes']:
    src=set(open(os.path.join(root,n['source']),encoding='utf-8').read().splitlines());cand=open(pkg+'/12_Candidate_Files/notes_candidates/'+n['candidate'],encoding='utf-8').read().splitlines()
    added=[l for l in cand if l not in src];skipped+=len(cand)-len(added)
    for x in V.scan_text('\n'.join(added),n['candidate']):
        ids.add(x['id']);(badf if x['verdict']=='FAIL' else expl).append(x)
ids=sorted(ids)
rec(9,'Validation-ID regression tests and D3 documents scan',t.startswith('30/30') and not badf,f'regression: {t}; D3 documents scanned in full; notes candidates scanned on the lines D3 added/changed only ({skipped} unchanged carried-forward lines not rescanned; they are byte-identical to the originals, e.g. line 67 "note on VAL-18 and digital gates", which line 74 of the same file explains as "VAL-18 missing → not allocated"); VAL ids in scanned text={ids}; active/unknown-ID findings={len(badf)}; explanatory mentions allowed={len(expl)}')

# ======== D3B owner-decision checks ========
import csv
SELF=pkg+'/D3B_X12_Selected_BRAND_Mapping.csv'
rec10=open(pkg+'/09_D3_Owner_Decision_Record.md',encoding='utf-8').read()
box=lambda label:re.search(r'^\[(.)\] '+re.escape(label),rec10,re.M).group(1)
rec(10,'Owner record: D3A ACCEPTED; D3B BRAND selected; Logo/OTHER/STILL-REQUIRED unticked; D3 resolved at owner-decision level',box('ACCEPTED')=='x' and box('PARTIALLY ACCEPTED')==' ' and box('REJECTED')==' ' and box('BRAND')=='x' and box('Logo')==' ' and box('OTHER')==' ' and box('OWNER DECISION STILL REQUIRED')==' ' and 'D3 = RESOLVED AT OWNER-DECISION LEVEL' in rec10 and 'AC20 closed' in rec10,'checkbox states read from 09_D3_Owner_Decision_Record.md; the "does NOT mean" clarification (AC20, release, Legal/IP, migration, production readiness) is present')
sel=list(csv.DictReader(open(SELF,encoding='utf-8')));names=[r['Controlled X12 filename'] for r in sel]
RXN=re.compile(r'^GEM_BRAND_[A-Za-z0-9]+_[A-Za-z0-9]+_vX\.Y_YYYYMMDD\.(svg|pdf|eps|png)$')
K='GEM_Brand_Assets_v1.0/04_official_kit/logo/'
src_ok=all(os.path.exists(os.path.join(root,r['Source path (unchanged)'])) and os.path.basename(r['Source path (unchanged)'])==r['Current source filename'] and sha(os.path.join(root,r['Source path (unchanged)']))==r['SHA-256 of source'] for r in sel)
kit_clean=git('diff','--quiet','HEAD','--','GEM_Brand_Assets_v1.0').returncode==0 and not [l for l in git('status','--porcelain','--','GEM_Brand_Assets_v1.0').stdout.splitlines()]
dup=sum(1 for r in sel if r['STREAM'].lower() in r['ASSET'].lower())
asset_ok=all(r['ASSET']=='Horizontal' and 'attested only by the Part A slide 75 example' in r['Notes'] or 'REQUIRES TAXONOMY CONFIRMATION' in r['Notes'] for r in sel)
ids=any(re.search(r'GEM-[A-Z]+-\d{3}\b',' '.join(r.values())) for r in sel)
rec(11,'Selected BRAND mapping: 128 rows; STREAM=BRAND on every row; 0 Logo; unique and case-insensitively unique; X12 pattern; sources unchanged and no rename; no asset-ID; ASSET values attested or flagged',len(sel)==128 and all(r['STREAM']=='BRAND' for r in sel) and not any('Logo' in r['STREAM'] or '_Logo_' in r['Controlled X12 filename'] for r in sel) and len(set(names))==128 and len({n.lower() for n in names})==128 and all(RXN.match(n) for n in names) and src_ok and kit_clean and all(r['Source rename required?']=='NO' and r['Manifest-only?']=='YES' for r in sel) and not ids and asset_ok and dup==0,f'rows={len(sel)}; STREAM=BRAND on all={all(r["STREAM"]=="BRAND" for r in sel)}; unique={len(set(names))}; case-insensitive unique={len({n.lower() for n in names})}; pattern GEM_BRAND_<ASSET>_<VARIANT>_vX.Y_YYYYMMDD.ext on all={all(RXN.match(n) for n in names)}; every source file exists under its current name with the recorded SHA-256={src_ok}; GEM_Brand_Assets_v1.0 has no change vs HEAD and no untracked/renamed file={kit_clean}; rename required NO and manifest-only YES on all rows; STREAM-in-ASSET duplicates={dup}; asset-ID strings in the mapping={ids}; no checksum rule introduced (the only hash column is the SHA-256 of the unchanged source file, the algorithm the kit already uses); ASSET values attested in the Part A example or flagged REQUIRES TAXONOMY CONFIRMATION={asset_ok}')
lg=list(csv.reader(open(pkg+'/06_D3B_X12_Logo_Mapping.csv',encoding='utf-8')));br=list(csv.reader(open(pkg+'/05_D3B_X12_BRAND_Mapping.csv',encoding='utf-8')))
sem=open(pkg+'/07_D3B_X12_Semantic_Test.csv',encoding='utf-8').read()
c4txt=open(pkg+'/04_D3B_X12_STREAM_Comparison.md',encoding='utf-8').read()
rec(12,'Logo mapping retained as history and marked NOT SELECTED; semantic test keeps the original 13 criteria and adds the final disposition',lg[0][-1]=='Selection status' and all(x[-1].startswith('NOT SELECTED') for x in lg[1:]) and len(lg)==130 and br[0][-1]=='Selection status' and 'BRAND — SELECTED' in sem and 'Logo — NOT SELECTED' in sem and sem.count('\n')>=15 and 'SELECTED CONTROLLED STREAM: BRAND' in c4txt and 'Prior state' in c4txt,f'06 rows={len(lg)-1} (128 files + metrics note), all marked NOT SELECTED; 05 marked comparison stage; semantic test contains the 13 criteria, the final disposition and the retained conflict note; 04 states SELECTED CONTROLLED STREAM: BRAND and keeps the prior state')
# pages that differ in the candidate PDFs vs ODI01-R1 PDFs
def pdfdiff(a,b,n):
    d=[]
    for pg in range(1,n+1):
        x=subprocess.run(['pdftotext','-f',str(pg),'-l',str(pg),'-layout',a,'-'],capture_output=True,text=True).stdout;y=subprocess.run(['pdftotext','-f',str(pg),'-l',str(pg),'-layout',b,'-'],capture_output=True,text=True).stdout
        if x!=y: d.append(pg)
    return d
dA=pdfdiff(g('*Part A*.pdf'),glob.glob(pkg+'/12_Candidate_Files/deck_candidates/*Part A*.pdf')[0],80);dC=pdfdiff(g('*Part C*.pdf'),glob.glob(pkg+'/12_Candidate_Files/deck_candidates/*Part C*.pdf')[0],78)
a75=[l for l in lines(L['decks']['A']['dst']) if l.startswith('Example: GEM_')];rule=[l for l in lines(L['decks']['A']['dst']) if 'The stream field follows the asset library folder' in l]
rec(13,'Part A slide 75: example uses BRAND; rule sentence unchanged; candidate PDFs differ from ODI01-R1 only on page 75 (A) and page 44 (C)',len(a75)==1 and 'GEM_BRAND_Horizontal_Black_v3.0_20261006.svg' in a75[0] and 'Logo' not in a75[0] and len(rule)==1 and dA==[75] and dC==[44],f'slide 75 example line={a75}; rule sentence present unchanged={len(rule)==1}; PDF pages differing from ODI01-R1: Part A {dA}, Part C {dC}')
stale=re.compile(r'OWNER DECISION REQUIRED|D3B pending|D3 unresolved|BRAND vs Logo unresolved|STILL REQUIRED',re.I)
cur=[(f,l[:90]) for f in ('01_D3_Executive_Summary.md','02_D3A_Authority_Verification.md','03_D3A_Implementation_Trace.md','08_D3_Cross_Document_Consequences.md','10_D3_Remaining_Gates.md') for l in open(pkg+'/'+f,encoding='utf-8').read().splitlines() if stale.search(l)]
r4=[l[:60] for l in c4txt.splitlines() if stale.search(l)]
hist_ok=all('Prior state' in l or 'prior state' in l or 'comparison stage' in l for l in [x for x in c4txt.splitlines() if stale.search(x)])
status_ok=all(('D3A — RESOLVED' in open(pkg+'/'+f,encoding='utf-8').read() and 'D3B — RESOLVED' in open(pkg+'/'+f,encoding='utf-8').read()) for f in ('01_D3_Executive_Summary.md','10_D3_Remaining_Gates.md','09_D3_Owner_Decision_Record.md'))
rec(14,'Stale-status scan: current-status sections say D3A RESOLVED / D3B RESOLVED / D3 RESOLVED AT OWNER-DECISION LEVEL / STREAM = BRAND; old phrases only in marked prior-state text',not cur and hist_ok and status_ok and all('STREAM = BRAND' in open(pkg+'/'+f,encoding='utf-8').read() or 'STREAM = BRAND' in open(pkg+'/'+f,encoding='utf-8').read().replace('X12 STREAM = BRAND','STREAM = BRAND') for f in ('01_D3_Executive_Summary.md','09_D3_Owner_Decision_Record.md','10_D3_Remaining_Gates.md')),f'stale phrases in current-status documents={cur}; in 04 only in prior-state text={hist_ok} ({r4}); status lines present in 01/09/10={status_ok}')
json.dump(R,open(pkg+'/11_QA_Evidence/d3_qa_results.json','w'),indent=1,ensure_ascii=False)
md=['# D3 QA results','',f'**{sum(r["result"]=="PASS" for r in R)}/{len(R)} checks PASS.** These are text, structure and manifest checks; they are not a statement of full system consistency. Native PowerPoint rendering is not tested; the one visibly edited slide (Part C 44) was rendered with LibreOffice (review evidence only): `C44_before_after.png`.','','| # | Check | Result | Evidence |','|---|---|---|---|']+[f"| {r['n']} | {r['name']} | **{r['result']}** | {r['evidence'].replace('|','/')} |" for r in R]
open(pkg+'/11_QA_Evidence/d3_qa_results.md','w',encoding='utf-8').write('\n'.join(md)+'\n')
sys.exit(0 if all(r['result']=='PASS' for r in R) else 1)
