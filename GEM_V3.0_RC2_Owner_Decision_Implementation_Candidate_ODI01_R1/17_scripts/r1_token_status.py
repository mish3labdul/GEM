"""ODI01-R1 D4 + D5 token implementation (R1: D5 note now states FUTURE DESIGN INTENT = 500 / CURRENT IMPLEMENTATION = 400; D4 logic and verifications unchanged from ODI01).
D4: accept all AFC02 downward status normalizations (weakest-dependency rule). Values/names/meaning untouched; nothing promoted.
D5: annotate the four weight-500 tokens "500 WEIGHT PENDING ACCEPTED FONT FILE — use 400 until accepted" (description/comment only).
Usage: token_status_implementation.py <repo_root> <out_dir>"""
import json,csv,sys,re,os,copy,collections
from pptx import Presentation
root,out=sys.argv[1],sys.argv[2]
SRC=os.path.join(root,'PartB_RC2/05_release/tokens/gem-tokens.v3.0-rc2')
d=json.load(open(SRC+'.json'))
RANK={'APPROVED':3,'CONDITIONAL':2,'PENDING VALIDATION':1}; NAME={v:k for k,v in RANK.items()}
def walk(n,p=()):
    if isinstance(n,dict):
        if '$value' in n: yield '.'.join(p),n
        else:
            for k,v in n.items():
                if not k.startswith('$'): yield from walk(v,p+(k,))
T=dict(walk(d))
EXT={'ARABIC_HIERARCHY':('PENDING VALIDATION','Part B slide 12 "ARABIC AND RTL · PENDING VALIDATION"; Register M02/M04; VAL-07'),
 'FILE_500':('CONDITIONAL','No accepted weight-500 font file in main; Register Y03 Evidence Required; VAL-05; VAL-19'),
 'FILE_500_ARABIC':('PENDING VALIDATION','As FILE_500 and Noto Sans Arabic is interim (M02, VAL-07)')}
def deps(n,t):
    D=[];parts=n.split('.')
    if n.startswith('type.lineHeight.'):
        s=T.get('type.size.'+parts[-1])
        if s: D.append(('type.size.'+parts[-1],s['status'],'Part B type table size/line pair; Register L09'))
    if n.startswith('type.') and re.search('arabic',n,re.I) and not n.startswith(('type.tracking.arabic','type.family.arabic')) and n!='type.arabicLineHeightMin':
        D+= [('type.family.arabic',T['type.family.arabic']['status'],EXT['ARABIC_HIERARCHY'][1]),('type.arabicLineHeightMin',T['type.arabicLineHeightMin']['status'],'Part B Arabic hierarchy, M04'),('ARABIC_HIERARCHY',)+EXT['ARABIC_HIERARCHY']]
    if n in('type.weight.medium','type.styleWeight.label'): D.append(('FILE_500',)+EXT['FILE_500'])
    if n.startswith('type.styleWeight.arabic'): D.append(('FILE_500_ARABIC',)+EXT['FILE_500_ARABIC'])
    v=t['$value']
    if isinstance(v,str):
        for ref in re.findall(r'\{([^}]+)\}',v):
            r=T.get(ref); D.append((ref,r['status'],'Token alias reference') if r else (ref,'UNRESOLVED','Alias target not found'))
    return D
W500=['type.weight.medium','type.styleWeight.label','type.styleWeight.arabicH3','type.styleWeight.arabicLabel']
NOTE='FUTURE DESIGN INTENT = 500 · CURRENT IMPLEMENTATION = 400 · 500 WEIGHT PENDING ACCEPTED FONT FILE (D5; Y03, VAL-05, VAL-19)'
cand=copy.deepcopy(d);CT=dict(walk(cand));rows=[];changed=[]
for n,t in T.items():
    D=deps(n,t);cur=t['status']
    if D:
        ws=min(RANK.get(s,0) for _,s,_ in D);wk=NAME.get(ws,'UNRESOLVED');new=NAME[min(RANK[cur],ws)] if ws else cur
        dep='; '.join(f'{a} [{b}]' for a,b,_ in D);auth='; '.join(sorted(set(c for *_,c in D)))
    else: wk='— (no dependency edge)';new=cur;dep='—';auth=f'Register {t["decision"]}'
    if new!=cur: CT[n]['status']=new;changed.append(n)
    if n in W500:
        CT[n]['description']=(t.get('description','')+' · ' if t.get('description') else '')+NOTE
    rows.append(dict(Token=n,orig=cur,dep=dep,wk=wk,new=new,auth=auth+(' | D4 owner decision 2026-10-08' if new!=cur else ''),n500=n in W500))
cand['meta']=dict(cand['meta']);cand['meta']['status']='ODI01-R1 CANDIDATE · D4 STATUS NORMALIZATION + D5 500-WEIGHT ANNOTATION · UNAPPROVED · NOT RELEASED (AC20 OPEN)'
cand['meta']['note']+=' ODI01-R1 candidate: statuses normalized by the weakest-dependency rule (owner decision D4); numeric values unchanged; weight-500 tokens annotated as future design intent (D5).'
os.makedirs(out,exist_ok=True)
json.dump(cand,open(out+'/gem-tokens.v3.0-rc2.ODI01-R1-candidate.json','w'),indent=2,ensure_ascii=False)
css=open(SRC+'.css',encoding='utf-8').read()
def cssname(n): return '--gem-'+re.sub(r'(?<=[a-z0-9])([A-Z])',lambda m:'-'+m.group(1).lower(),n.replace('.','-')).lower()
cnt=0
for n in changed:
    css,k=re.subn(r'(%s:[^;]*;\s*/\*\s*)(APPROVED|CONDITIONAL|PENDING VALIDATION)'%re.escape(cssname(n)),lambda m:m.group(1)+CT[n]['status'],css);cnt+=k
c5=0
for n in W500:
    css,k=re.subn(r'(%s:[^;]*;\s*/\*[^*]*?)(\s*\*/)'%re.escape(cssname(n)),lambda m:m.group(1)+' · FUTURE DESIGN INTENT = 500 · CURRENT IMPLEMENTATION = 400 · 500 WEIGHT PENDING ACCEPTED FONT FILE'+m.group(2),css,count=1);c5+=k
css=css.replace('WORKING SPECIFICATION · NOT RELEASED (AC20 PENDING).','WORKING SPECIFICATION · NOT RELEASED (AC20 PENDING). ODI01-R1 CANDIDATE · statuses normalized (D4) · 500 = future design intent, current implementation 400 (D5) · UNAPPROVED.',1)
open(out+'/gem-tokens.v3.0-rc2.ODI01-R1-candidate.css','w',encoding='utf-8').write(css)
# ---------- verifications ----------
V=collections.OrderedDict()
O=dict(walk(d));C=dict(walk(cand))
V['Value parity (value+type+name set, 153 tokens)']= 'PASS' if {k:(v['$value'],v['$type']) for k,v in O.items()}=={k:(v['$value'],v['$type']) for k,v in C.items()} and len(C)==153 else 'FAIL'
V['Decision IDs unchanged']='PASS' if all(O[k]['decision']==C[k]['decision'] for k in O) else 'FAIL'
V['Only status/description changed (no other key touched)']='PASS' if all({a:b for a,b in O[k].items() if a not in('status','description')}=={a:b for a,b in C[k].items() if a not in('status','description')} for k in O) else 'FAIL'
V['Description changed only on the 4 weight-500 tokens']='PASS' if [k for k in O if O[k].get('description')!=C[k].get('description')]==W500 else 'FAIL'
V['No promotion (every status ≤ original)']='PASS' if all(RANK[C[k]['status']]<=RANK[O[k]['status']] for k in O) else 'FAIL'
nchg=[k for k in O if O[k]['status']!=C[k]['status']]
V['26 normalizations applied']='PASS' if len(nchg)==26 else f'FAIL ({len(nchg)})'
tr=collections.Counter((O[k]['status'],C[k]['status']) for k in nchg)
# JSON/CSS parity
decl_o=re.findall(r'(--gem-[A-Za-z0-9-]+)\s*:\s*([^;]+);',open(SRC+'.css',encoding='utf-8').read());decl_c=re.findall(r'(--gem-[A-Za-z0-9-]+)\s*:\s*([^;]+);',css)
V['CSS declarations identical (name+value)']='PASS' if decl_o==decl_c else 'FAIL'
cstat=dict((m.group(1),m.group(2)) for m in re.finditer(r'(--gem-[A-Za-z0-9-]+)\s*:[^;]*;\s*/\*\s*(APPROVED|CONDITIONAL|PENDING VALIDATION|REFERENCE)',css))
mis=[k for k in C if cssname(k) in cstat and cstat[cssname(k)]!=C[k]['status']]
V[f'JSON↔CSS status parity ({len([k for k in C if cssname(k) in cstat])} tokens with a CSS status comment; {len(C)-len([k for k in C if cssname(k) in cstat])} without)']='PASS' if not mis else 'FAIL '+str(mis[:3])
V['CSS status comments updated = 26']='PASS' if cnt==26 else f'FAIL ({cnt})'
V['D5 CSS annotations on 4 weight-500 tokens']='PASS' if c5==4 else f'FAIL ({c5})'
# dependency hierarchy re-run on candidate
bad=[]
T2=C
for k,t in C.items():
    for a,s,_ in deps(k,t):
        if RANK.get(s,0) and RANK[t['status']]>RANK[s] : bad.append(k)
        if a in T2 and RANK[T2[a]['status']]<RANK[t['status']] and a!=k: bad.append(k)
V['Dependency hierarchy: unexplained inflation (candidate)']='PASS — 0' if not bad else 'FAIL '+str(sorted(set(bad))[:5])
# Arabic
ar=[k for k in C if k.startswith('type.') and re.search('arabic',k,re.I) and not k.startswith(('type.tracking.arabic','type.family.arabic')) and k!='type.arabicLineHeightMin']
V[f'Arabic hierarchy tokens not stronger than PENDING VALIDATION ({len(ar)} tokens)']='PASS' if all(C[k]['status']=='PENDING VALIDATION' for k in ar) else 'FAIL'
# non-Arabic conditional
lh=[k for k in C if k.startswith('type.lineHeight.') and 'arabic' not in k.lower()]
V['Non-Arabic lineHeight ≤ paired size status']='PASS' if all(RANK[C[k]['status']]<=RANK[C['type.size.'+k.split('.')[-1]]['status']] for k in lh) else 'FAIL'
V['Non-Arabic CONDITIONAL size → lineHeight ≤ CONDITIONAL']='PASS' if all(C[k]['status']!='APPROVED' for k in lh if C['type.size.'+k.split('.')[-1]]['status']=='CONDITIONAL') else 'FAIL'
# Part B table parity: Part B tables carrying status words
deck=Presentation(f'{root}/PartB_RC2/05_release/GEM Digital Design System V3.0 — Part B — RC2.pptx')
t21=[[c.text.strip() for c in r.cells] for sh in deck.slides[20].shapes if getattr(sh,'has_table',False) and sh.has_table for r in sh.table.rows][1:]
mp={'locale-source':'locale.source','locale-arabic':'locale.arabic','numerals-technical':'locale.numeralsTechnical','numerals-guestArabic':'locale.numeralsGuestArabic','calendar-default':'locale.calendarDefault'}
okp=True;detail=[]
for nm,val,st in t21:
    if nm in mp:
        tk=C[mp[nm]];good=(tk['status']==st and str(tk['$value'])==val);okp&=good;detail.append(f'{nm}:{st}={tk["status"]}')
V['Part B slide 21 locale table ↔ candidate tokens (value+status, 5 rows)']='PASS' if okp else 'FAIL '+'; '.join(detail)
V['Part B slides 10/12 typography tables: carry no status column → size/line/weight/tracking values re-verified by audit_tokens.py']='SEE audit_tokens.py (run separately)'
# CSV
H=['Token','Original status','Dependency','Weakest dependency status','New status','Value changed?','Authority','QA result']
with open(out+'/token_status_implementation.csv','w',newline='',encoding='utf-8') as f:
    w=csv.writer(f);w.writerow(H)
    for r in rows:
        k=r['Token'];vch='NO' if O[k]['$value']==C[k]['$value'] and O[k]['$type']==C[k]['$type'] else 'YES'
        q='PASS' if vch=='NO' and RANK[r['new']]<=RANK[r['orig']] and (r['new']==C[k]['status']) else 'FAIL'
        w.writerow([k,r['orig'],r['dep'],r['wk'],r['new'],vch,r['auth']+(' | D5: weight-500 annotation only' if r['n500'] else ''),q])
res={'verifications':V,'transitions':{f'{a} -> {b}':n for (a,b),n in tr.items()},'changed':nchg,'n500':W500}
json.dump(res,open(out+'/token_verification.json','w'),indent=1)
print(json.dumps(res,indent=1))
