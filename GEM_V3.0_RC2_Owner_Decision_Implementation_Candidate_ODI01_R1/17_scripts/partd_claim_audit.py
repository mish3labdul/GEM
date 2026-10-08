"""ODI01 Part D claim audit (classes: BRAND RULE APPROVED / OWNER-DESIGNATED CONCEPT / EVIDENCE REQUIRED / PRODUCTION STATUS / UNSUPPORTED CLAIM). Reads the Part D PPTX (shapes, tables, groups, notes) and classifies every line containing a claim keyword.
Usage: partd_claim_audit.py <repo_root> <out_csv> [pptx_path]"""
import sys,glob,re,csv,collections,os
from pptx import Presentation
root,out=sys.argv[1],sys.argv[2]
f=sys.argv[3] if len(sys.argv)>3 else glob.glob(root+'/Amenities_Portfolio_PartD_RC2/04_release/*.pptx')[0]
has_registers=os.path.exists(root+'/Amenities_Portfolio_PartD_RC2/01_mockups/MOCKUP_REGISTER.csv')
KW=re.compile(r'\b(approved|approval|accepted|validated|validation|final|production|owner-approved|owner|proof|master|chat product direction|supporting csv registers)\b',re.I)
def classify(t):
    L=t.lower()
    if 'supporting csv registers' in L: return 'UNSUPPORTED CLAIM','R-dangling','Notes cite supporting CSV registers (concept, source, asset, change, open decisions). Only 01_mockups/MOCKUP_REGISTER.csv exists in the package.'
    if 'chat product direction' in L: return 'OWNER-DESIGNATED CONCEPT','R-owner','Source cited as "chat product direction": an owner/product direction, not a validated rule; unverifiable in the repository.'
    if 'once supplied and approved' in L: return 'EVIDENCE REQUIRED','R-future','Dieline is conditional on supplier supply and approval (ODI01 D6 wording).'
    if 'a mockup register is provided' in L: return 'OWNER-DESIGNATED CONCEPT','R-register','ODI01 D6 wording: states only what the package contains (MOCKUP_REGISTER.csv).'
    if 'approved supplier dieline' in L: return 'EVIDENCE REQUIRED','R-future','Reads as if an approved dieline exists; none does (supplier proof gate). Candidate wording provided.'
    if 'approved descriptive stream' in L: return 'BRAND RULE APPROVED','R-brand','Register T01/AC03 (architecture); capability evidence stays conditional (same note).'
    if 'six approved hierarchy levels' in L or 'register’s approved order' in L or "register's approved order" in L: return 'BRAND RULE APPROVED','R-brand','Six-level hierarchy is a Register decision (T05).'
    if L.startswith('authority:'): return 'OWNER-DESIGNATED CONCEPT','R-disclaimer','Boilerplate that negates release/production authorization and states the concept status.'
    if 'concept / not production' in L or 'structural concept' in L or 'working edition / concept' in L: return 'OWNER-DESIGNATED CONCEPT','R-label','Required concept marking retained.'
    if re.search(r'proof required|physical proof|pending physical proof|validation pending|production validation|pending validation|localization approval|requires owner approval|require approval|require supplier|remain open|pending production|production master acceptance pending|touchpoint scope conditional|release authorization|physical validation',L): return 'PRODUCTION STATUS' if re.search(r'production validation|pending production|production master acceptance|release authorization',L) else 'EVIDENCE REQUIRED','R-gate','States an open gate.'
    if re.search(r'not an approved production asset|no item in this portfolio|approved logo reproduction|no approval or capability claim|not .*approved master ids',L): return 'OWNER-DESIGNATED CONCEPT','R-negation','Negative statement: no approval claimed.'
    if L in('gem master artwork',): return 'EVIDENCE REQUIRED','R-master','"Master artwork" label while production-master acceptance is open (AC07 Evidence Required, VAL-02).'
    if re.search(r'one master brand|master logo|master identity',L): return 'BRAND RULE APPROVED','R-brand','Master-brand architecture (Register, Part A).'
    if re.search(r'final approval|correct master|^production$|production release|^release$|^proof$|master assets|production specification|supplier proof|production$|^step 11|no step is bypassed|validation dependency|^proof',L): return 'PRODUCTION STATUS','R-process','Describes the Part C release sequence; no claim about any item.'
    return 'OWNER-DESIGNATED CONCEPT','R-default','Keyword appears in a process or label context; no approval claim.'
P=Presentation(f); seen=collections.OrderedDict()
for i,s in enumerate(P.slides,1):
    parts=[]
    def add(sh,k):
        if sh.has_text_frame: parts.append((k,sh.text_frame.text))
        if getattr(sh,'has_table',False) and sh.has_table: parts.extend(('table',c.text) for r in sh.table.rows for c in r.cells)
        if sh.shape_type==6:
            for g in sh.shapes: add(g,'group')
    for sh in s.shapes: add(sh,'shape')
    if s.has_notes_slide: parts.append(('notes',s.notes_slide.notes_text_frame.text))
    for k,t in parts:
        for line in t.split('\n'):
            line=line.strip()
            if line and KW.search(line):
                # split long notes into sentences so each claim is judged alone
                units=[line] if k!='notes' else [u.strip() for u in re.split(r'(?<=[.!?])\s+(?=[A-Z])',line) if KW.search(u)]
                for u in units: seen.setdefault((k,u),[]).append(i)
rows=[]
for (k,t),sl in seen.items():
    c,rule,why=classify(t); rows.append([c,k,','.join(map(str,sorted(set(sl)))),len(sl),t,rule,why])
rows.sort(key=lambda r:(['BRAND RULE APPROVED','OWNER-DESIGNATED CONCEPT','EVIDENCE REQUIRED','PRODUCTION STATUS','UNSUPPORTED CLAIM'].index(r[0]),r[2]))
H=['Classification','Where','Slides','Occurrences','Text','Rule','Reason']
csv.writer(open(out,'w',newline='',encoding='utf-8')).writerows([H]+rows)
cnt=collections.Counter(); 
for r in rows: cnt[r[0]]+=r[3]
print(len(rows),'unique claims;',sum(r[3] for r in rows),'occurrences',dict(cnt))
