from pathlib import Path
from docx import Document
from docx.shared import Pt,Mm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path('work/GEM/GEM_Letterhead_Set_v1.1');d=Document();s=d.sections[0]
s.page_width=Mm(210);s.page_height=Mm(297);s.left_margin=s.right_margin=Mm(23);s.top_margin=Mm(23);s.bottom_margin=Mm(25);s.footer_distance=Mm(12)
for name,size,font in [('Normal',10,'Inter'),('Title',21,'Jost'),('Heading 1',12,'Jost')]:
 st=d.styles[name];st.font.name=font;st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('12171D');st.font.bold=False
 st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.35
for e in list(d.styles.element.xpath('.//w:pBdr')):e.getparent().remove(e)
def p(t,style=None):return d.add_paragraph(t,style)
def h(t):p(t,'Heading 1')
def source(t):
 pp=p('Source: '+t)
 for r in pp.runs:r.font.size=Pt(9)
A='GEM_V3_RC2/04_release/GEM Brand Guidelines V3.0 — Part A — RC2.pptx (23, 34–36, 67)'
B='PartB_RC2/05_release/GEM Digital Design System V3.0 — Part B — RC2.pptx (9–13); PartB_RC2/05_release/tokens/gem-tokens.v3.0-rc2.json'
C='GEM_V3_RC2/04_release/GEM Production Standards V3.0 — Part C — RC2.pptx (5, 19–21, 44)'
REG='GEM_V3_RC2/00_originals/GEM_V3_Final_Brand_Approval_Register_Prefilled.xlsx'
p('GEM Letterhead Specification','Title');p('Set v1.1 Application Revision 02 • 7 October 2026\nWORKING APPLICATION / PENDING VALIDATION')
p('This sheet records the corrected RC2 application settings. Approved brand rules remain separate from proposed stationery settings and pending production values. It does not authorize release or manufacture.')
h('A4 page and writing area')
p('APPLICATION SETTING: A4 portrait, 210 × 297 mm. Side margins 25 mm. Top 30 mm, Executive 36 mm. Bottom 37 mm. Header distance 18 mm and footer distance 12 mm. Writing width 160 mm; English running body uses a 30 mm end indent for a 130 mm measure. Arabic uses 160 mm. First-page body starts below the logo/tagline and varies with content. Continuation starts at the 30 mm top margin.')
p('Footer is a four-line maximum zone: address, contact/page fields, working status and (where relevant) localization status. Essential footer text is Inter Regular 9 pt with 12.6 pt exact leading. No geometry crosses the writing area. Mechanical print-safe zone remains REQUIRES SUPPLIER; these margins are application allowances.')
source(A+'; '+C)
h('Supplied artwork and clearspace')
p('APPROVED BRAND RULE: supplied vector artwork only, unchanged proportion, colour and orientation. Official kit acceptance as production master remains pending.')
p('APPLICATION SETTING: GEM_Brand_Assets_v1.0/04_official_kit/logo/horizontal/svg/gem-horizontal-ink.svg at 32 mm; Executive 36 mm. Minimal uses gem-horizontal-black.svg at 32 mm. DOCX carries the unchanged supplied 2048 px PNG compatibility fallback. PDF uses vector paths. At 32 mm, 1u = 5.19 mm and preferred 2u = 10.39 mm; Executive 2u = 11.69 mm. Tagline gap 12 mm. No new symbol or spark is added.')
source('GEM_Brand_Assets_v1.0/README.md; 04_official_kit/logo/metrics.json; '+REG+' H10–H14')
p('Type and language','Title').paragraph_format.page_break_before=True
h('Typography and editing')
p('APPROVED BRAND RULE: Jost for display/subject; Inter for office body; Noto Sans Arabic for working Arabic subject to localization/font acceptance. Zero body/Arabic tracking. Tagline “HOSPITALITY, IN PERFECT PROPORTION” uses 0.18em tracking.')
p('APPLICATION SETTING: Inter Regular body 10.5 pt / 16.8 pt exact leading. English subject real Heading 1, Jost Regular 12 pt. Arabic body Noto Sans Arabic Regular 11.5 pt / 20.125 pt; subject 13 pt. Metadata 9 pt Inter, Arabic labels 10 pt. Tagline Jost Regular 10.5 pt with 38 twips (1.9 pt) expanded spacing, Word approximation of 0.18em. No bold synthesis. Signature reserve 16 mm. Essential footer/status text 9 pt; digital legibility decision remains REQUIRES OWNER, physical minimum remains PENDING PHYSICAL PROOF.')
source(B+'; '+REG+' B03, L01–L09')
h('Arabic and bilingual flow')
p('Arabic paragraphs use genuine RTL flow. Complete Latin references, dates, email, phone and URL fields remain Inter LTR, unmirrored and untracked. Date/reference are separate lines to preserve usable identifiers. Arabic placeholder text is not approved legal wording. Every Arabic/bilingual footer retains PENDING LOCALIZATION APPROVAL.')
p('Bilingual content is sequential: Arabic recipient/subject/body first, then English equivalent, followed by signatory details. English/transliteration signatory field is optional. S07 explicitly approves Saudi/GCC wayfinding hierarchy; applying that hierarchy to correspondence remains PROPOSED APPLICATION / REQUIRES OWNER. Latin technical identifiers retain Latin digits; business numerals/date policy requires contextual localization approval.')
source(B+'; '+REG+' S07, M01–M12; qa/PartB_RC2_RTL_Localization_QA_Report.md')
h('Continuation and field replacement')
p('Start a complete letter in a First Page file; subsequent pages automatically use the quiet identifier and footer. Continuation is a companion. Replace SAMPLE CONTENT and all body, header and footer fields before sending. Update live PAGE/NUMPAGES in Word before export. Allow long copy to repaginate.')
p('Output and release control','Title').paragraph_format.page_break_before=True
h('Digital PDF and accessibility')
p('Selectable text, font embedding, A4 geometry and semantic paragraph/heading structure are checked in the evidence register. Exporter omissions in combining-mark ToUnicode maps are repaired from its explicit ActualText clusters without altering visible glyphs. Logical source paragraphs are retained in semantic ActualText. Some generic extractors still reverse or omit mixed RTL fragments; native copy/paste, reading order and assistive-technology acceptance remain PENDING. Re-exporting a DOCX does not automatically include these PDF repairs.')
p('Use the supplied digital PDF as a sample portfolio, not a blank correspondence form. Export and validate every completed real letter. Links require verified contact values; no real email, telephone or URL is invented. Real headings, run language tags and meaningful supplied-logo descriptions are present. Tagged structure is not a claim of complete accessibility validation.')
source(REG+' V02, V05, V13; '+B+'; qa/PartB_RC2_Accessibility_QA_Report.md')
h('Print and grayscale')
p('APPROVED BRAND RULE: INK #12171D, BEIGE #BCACA7, BLACK #020202, WHITE #FFFFFF. Ink/White contrast approximately 18:1; Black/White 20.8:1. Beige is decorative rule only. Minimal uses the supplied Black logo without tagline or footer rule. No background flood, bleed artwork, gold, gradient or added contact icons.')
p('PENDING PRODUCTION VALUE: CMYK/Pantone, output profile/intent, PDF/X, paper stock/GSM, finishing, tolerances and physical minimum type/logo sizes. Print PDF is the same A4 RGB artwork with print-candidate metadata; it is not supplier-ready commercial production artwork. REQUIRES SUPPLIER, PENDING PREPRESS STANDARD and PENDING PHYSICAL PROOF remain open.')
source(C+'; '+REG+' AC20; qa/GEM_V3_RC2_Final_Open_Evidence_Register.md')
h('Provenance and approval')
p('Original v1.0 and its conflicting release derivatives are preserved. This revision uses 07_final_audit/corrected_docx as its working baseline, retaining verified useful release edits. Source snapshot, reconciliation, exact assets, font hashes and changes are included. No External Partners Brief is used. Brand Owner, Arabic/localization, asset/font/license, accessibility and supplier gates remain open. Author attribution is unset because it is unverified.')
f=s.footer.paragraphs[0];f.text='GEM™ • Set v1.1 Application Revision 02 • WORKING APPLICATION / PENDING VALIDATION'
for r in f.runs:r.font.name='Inter';r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('12171D')
d.core_properties.title='GEM Letterhead Specification Set v1.1 Application Revision 02';d.core_properties.author='';d.core_properties.language='en-GB';d.core_properties.subject='WORKING APPLICATION / PENDING VALIDATION'
# Remove theme-font inheritance so Office-compatible exports cannot substitute Carlito.
for st in d.styles:
 if hasattr(st.element,'get_or_add_rPr'):
  rp=st.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
  if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
  for key in list(rf.attrib):del rf.attrib[key]
  family='Jost' if st.name in ['Title','Heading 1'] else 'Inter'
  for key in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+key),family)
for pp in list(d.paragraphs)+list(s.footer.paragraphs):
 for rr in pp.runs:
  family='Jost' if pp.style.name in ['Title','Heading 1'] else 'Inter';rr.font.name=family;rr.font.bold=False
  rf=rr._r.get_or_add_rPr().find(qn('w:rFonts'))
  for key in list(rf.attrib):del rf.attrib[key]
  for key in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+key),family)
d.save(R/'03_specs/GEM_Letterhead_Specification.docx')

from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
f=R/'03_specs/GEM_Letterhead_Specification.docx';base=ZipFile(R/'01_templates/GEM_Letterhead_English_First_Page.docx');data={n:ZipFile(f).read(n) for n in ZipFile(f).namelist()}
for n in base.namelist():
 if n.startswith('word/fonts/') or n in ['word/fontTable.xml','word/_rels/fontTable.xml.rels']:data[n]=base.read(n)
root=E.fromstring(data['[Content_Types].xml']);E.SubElement(root,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='odttf',ContentType='application/vnd.openxmlformats-officedocument.obfuscatedFont');data['[Content_Types].xml']=E.tostring(root,xml_declaration=True,encoding='UTF-8')
with ZipFile(f,'w',ZIP_DEFLATED) as z:
 for n,b in data.items():z.writestr(n,b)
