from pathlib import Path
import zipfile,uuid,hashlib,json,shutil,re
from lxml import etree as E
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent.parent; repo=R.parent
INK='12171D'; BEIGE='BCACA7'; BLACK='020202'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'; REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'; PKG='http://schemas.openxmlformats.org/package/2006/relationships'
FONTDIR=R/'00_source/fonts'
def el(name,**attrs):
 x=OxmlElement(name)
 for k,v in attrs.items():x.set(qn('w:'+k),str(v))
 return x

def run(p,text,font='Inter',size=10.5,bold=False,arabic=False,color=INK):
 r=p.add_run(text); r.font.name=font;r.font.size=Pt(size);r.font.bold=False;r.font.color.rgb=RGBColor.from_string(color)
 rp=r._element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
 for a in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+a),font)
 rp.append(el('w:lang',val='ar-SA' if arabic else 'en-GB',bidi='ar-SA' if arabic else 'en-GB'))
 rp.append(el('w:rtl',val='1' if arabic else '0'));rp.append(el('w:szCs',val=round(size*2)))
 rp.append(el('w:bCs',val='0'))
 return r

def para(d,text='',rtl=False,font=None,size=None,bold=False,after=6,before=0,style=None,keep=False):
 is_body=(not rtl and size is None and font is None and text.startswith(('Please share','We would','We request','This letter','Additional sample','SAMPLE CONTENT. This paragraph')))
 if is_body:style='Body Text'
 p=d.add_paragraph(style=style);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
 pp=p._p.get_or_add_pPr();pp.append(el('w:bidi',val='1' if rtl else '0'))
 if rtl:pp.find(qn('w:jc')).set(qn('w:val'),'start')
 p.paragraph_format.space_after=Pt(after);p.paragraph_format.space_before=Pt(before)
 p.paragraph_format.line_spacing=Pt((size or 11.5)*1.75) if rtl else Pt((size or 10.5)*1.6);p.paragraph_format.keep_with_next=keep;p.paragraph_format.widow_control=True
 if is_body:p.paragraph_format.right_indent=Mm(30)
 if text:run(p,text,font or ('Noto Sans Arabic' if rtl else 'Inter'),size or (11.5 if rtl else 10.5),bold,rtl)
 return p

def field(p,code):
 r=p.add_run(); f=el('w:fldSimple',instr=code);rr=el('w:r');pr=el('w:rPr');pr.append(el('w:rFonts',ascii='Inter',hAnsi='Inter',cs='Inter'));pr.append(el('w:sz',val=16));rr.append(pr);t=el('w:t');t.text='1';rr.append(t);f.append(rr);r._r.addnext(f)

def logo(p,width,color='ink'):
 path=repo/f'GEM_Brand_Assets_v1.0/04_official_kit/logo/horizontal/png/gem-horizontal-{color}-2048.png'
 if not path.exists():path=R/'00_source/reused_assets'/path.name
 pic=p.add_run().add_picture(str(path),width=Mm(width));pic._inline.docPr.set('descr','GEM trademark logo')
 # exact SVG source plus supplied PNG fallback
 svg=repo/f'GEM_Brand_Assets_v1.0/04_official_kit/logo/horizontal/svg/gem-horizontal-{color}.svg'
 if not svg.exists():svg=R/'00_source/reused_assets'/svg.name
 from docx.opc.part import Part
 from docx.opc.packuri import PackURI
 partname=PackURI('/word/media/gem-horizontal-'+color+'.svg')
 part=next((x for x in p.part.package.parts if x.partname==partname),None)
 if part is None:part=Part(partname,'image/svg+xml',svg.read_bytes(),p.part.package)
 rid=p.part.relate_to(part,REL+'/image')
 blip=pic._inline.xpath('.//a:blip')[0];extlst=E.SubElement(blip,'{http://schemas.openxmlformats.org/drawingml/2006/main}extLst')
 ext=E.SubElement(extlst,'{http://schemas.openxmlformats.org/drawingml/2006/main}ext',uri='{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
 E.SubElement(ext,'{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',{qn('r:embed'):rid})

def header(sec,rtl=False,cont=False,executive=False,minimal=False):
 def build(h,first):
  p=h.paragraphs[0];p._p.get_or_add_pPr().append(el('w:bidi',val=0));p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
  p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
  if first:
   width=36 if executive else 32;logo(p,width,'black' if minimal else 'ink')
   if not minimal:
    p=h.add_paragraph();p._p.get_or_add_pPr().append(el('w:bidi',val=0));p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before=Mm(12);p.paragraph_format.space_after=Mm(4)
    p.paragraph_format.line_spacing=Pt(15);r=run(p,'HOSPITALITY, IN PERFECT PROPORTION','Jost',10.5);r._element.get_or_add_rPr().append(el('w:spacing',val='38'))
   else:p.paragraph_format.space_after=Mm(12)
  else:
   run(p,'GEM™','Jost',9);run(p,'    [REFERENCE NUMBER]','Inter',8)
 if cont:
  sec.different_first_page_header_footer=False;build(sec.header,False)
 else:
  sec.different_first_page_header_footer=True;build(sec.first_page_header,True);build(sec.header,False)

def footer(sec,rtl=False,minimal=False):
 for f in ([sec.first_page_footer,sec.footer] if sec.different_first_page_header_footer else [sec.footer]):
  p=f.paragraphs[0];p._p.get_or_add_pPr().append(el('w:bidi',val=0));p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
  pp=p._p.get_or_add_pPr();b=el('w:pBdr');b.append(el('w:top',val='single',sz=4,color=BLACK if minimal else BEIGE,space=6));pp.append(b)
  p.paragraph_format.space_after=Pt(4);run(p,'GEM™   [ADDRESS] · [CITY] · [COUNTRY] · [POSTAL CODE]',size=7.5)
  p=f.add_paragraph();p._p.get_or_add_pPr().append(el('w:bidi',val=0));p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.space_after=Pt(6)
  run(p,'[T]   [E]   [W]',size=7.5);run(p,' '*5+'Page ',size=7.5);field(p,'PAGE');run(p,' of ',size=7.5);field(p,'NUMPAGES')
  p=f.add_paragraph();p._p.get_or_add_pPr().append(el('w:bidi',val=0));p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
  run(p,'WORKING APPLICATION / PENDING VALIDATION',size=6.5)

def configure_edit_styles(d,rtl=False):
 for name in ['Normal','Body Text','Heading 1','Heading 2','Heading 1 Char','Heading 2 Char']:
  if name not in d.styles:continue
  st=d.styles[name];style_rtl=rtl and name!='Body Text';rp=st.element.get_or_add_rPr()
  for child in list(rp):
   if child.tag in [qn('w:rFonts'),qn('w:b'),qn('w:bCs'),qn('w:lang'),qn('w:spacing'),qn('w:szCs')]:rp.remove(child)
  rp.append(el('w:rFonts',ascii='Inter' if name in ['Normal','Body Text'] else 'Jost',hAnsi='Inter' if name in ['Normal','Body Text'] else 'Jost',cs='Noto Sans Arabic',eastAsia='Inter'))
  rp.append(el('w:b',val='0'));rp.append(el('w:bCs',val='0'));rp.append(el('w:spacing',val='0'))
  rp.append(el('w:lang',val='ar-SA' if style_rtl else 'en-GB',bidi='ar-SA'))
  rp.append(el('w:szCs',val=23 if name in ['Normal','Body Text'] else 24))
  if st.type==1:
   pp=st.element.get_or_add_pPr();pp.append(el('w:bidi',val='1' if style_rtl else '0'));pp.append(el('w:jc',val='start' if style_rtl else 'left'))
   if name in ['Normal','Body Text']:
    pp.append(el('w:spacing',line=403 if style_rtl else 336,lineRule='exact',after=120))
   if name=='Body Text':st.paragraph_format.right_indent=Mm(30)
 defaults=d.styles.element.find(qn('w:docDefaults'));rp=defaults.find(qn('w:rPrDefault')).find(qn('w:rPr'))
 for child in list(rp):rp.remove(child)
 rp.append(el('w:rFonts',ascii='Inter',hAnsi='Inter',cs='Noto Sans Arabic',eastAsia='Inter'));rp.append(el('w:lang',val='ar-SA' if rtl else 'en-GB',bidi='ar-SA'));rp.append(el('w:sz',val=21));rp.append(el('w:szCs',val=23));rp.append(el('w:b',val=0));rp.append(el('w:bCs',val=0))

def embed_fonts(path):
 with zipfile.ZipFile(path) as z: data={n:z.read(n) for n in z.namelist()}
 ft=E.fromstring(data['word/fontTable.xml']);rels=E.Element('{'+PKG+'}Relationships');ct=E.fromstring(data['[Content_Types].xml'])
 E.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='odttf',ContentType='application/vnd.openxmlformats-officedocument.obfuscatedFont')
 for family,dirname in [('Inter','inter'),('Jost','jost'),('Noto Sans Arabic','notosansarabic')]:
  fs=ft.xpath('w:font[@w:name="'+family+'"]',namespaces={'w':W})
  node=fs[0] if fs else E.SubElement(ft,'{'+W+'}font',{'{'+W+'}name':family})
  for style in ['Regular']:
   font=FONTDIR/dirname/(family.replace(' ','')+'-'+style+'.ttf'); raw=bytearray(font.read_bytes());key=uuid.uuid4();kb=key.bytes[::-1]
   for i in range(32):raw[i]^=kb[i%16]
   filename=family.replace(' ','')+'-'+style+'.odttf';data['word/fonts/'+filename]=bytes(raw);rid='rId'+family.replace(' ','')+style
   E.SubElement(rels,'{'+PKG+'}Relationship',Id=rid,Type=REL+'/font',Target='fonts/'+filename)
   E.SubElement(node,'{'+W+'}embed'+style,{'{'+REL+'}id':rid,'{'+W+'}fontKey':'{'+str(key).upper()+'}','{'+W+'}subsetted':'false'})
 data['word/fontTable.xml']=E.tostring(ft,xml_declaration=True,encoding='UTF-8');data['word/_rels/fontTable.xml.rels']=E.tostring(rels,xml_declaration=True,encoding='UTF-8');data['[Content_Types].xml']=E.tostring(ct,xml_declaration=True,encoding='UTF-8')
 settings=E.fromstring(data['word/settings.xml']);settings.append(el('w:embedTrueTypeFonts'));settings.append(el('w:embedSystemFonts'));data['word/settings.xml']=E.tostring(settings,xml_declaration=True,encoding='UTF-8')
 with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
  for n,v in data.items():z.writestr(n,v)

def create(name,lang='en',cont=False,variant='',pages=1,stress=False,return_document=False):
 d=Document();sec=d.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
 sec.top_margin=Mm(30 if not variant=='Executive' else 36);sec.bottom_margin=Mm(37);sec.left_margin=sec.right_margin=Mm(25)
 sec.header_distance=Mm(18);sec.footer_distance=Mm(12)
 n=d.styles['Normal'];n.font.name='Inter';n.font.size=Pt(10.5);n.font.color.rgb=RGBColor.from_string(INK)
 for st in ['Heading 1','Heading 2']:d.styles[st].font.name='Jost';d.styles[st].font.size=Pt(12);d.styles[st].font.color.rgb=RGBColor.from_string(INK)
 rtl=lang!='en';configure_edit_styles(d,rtl);header(sec,rtl,cont,variant=='Executive',variant=='Minimal');footer(sec,rtl,variant=='Minimal')
 cp=d.core_properties;cp.title='GEM Branded Letterhead '+name;cp.subject='Working RC2 application pending validation';cp.author='GEM';cp.language='ar-SA' if rtl else 'en-GB'
 # Paragraph-level bidi governs correspondence; page section remains unmirrored
 p=para(d,'SAMPLE CONTENT'+(' · PENDING LOCALIZATION APPROVAL' if rtl else ''),size=7.5,after=9)
 if rtl:p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 if cont:
  para(d,'[Continuation text]' if lang=='en' else '[نص الصفحة التالية]',rtl,after=12)
 else:
  
  if rtl:
   p=para(d,'',True,after=8);run(p,'التاريخ: ',font='Noto Sans Arabic',size=10,arabic=True);run(p,'[DATE]',size=9);run(p,'    المرجع: ',font='Noto Sans Arabic',size=10,arabic=True);run(p,'[REFERENCE NUMBER]',size=9)
  else:para(d,'Date: [DATE]    Reference: [REFERENCE NUMBER]',size=8.5,after=10)
  if rtl:
   para(d,'«اسم المستلم»\n'+('«اسم الجهة الطويل لاختبار المراسلات الإدارية وطلبات المشتريات والمراجعة الداخلية»' if stress else '«اسم الجهة»')+'\n«العنوان»',True,after=5,keep=True)
   if lang=='bi':para(d,'[RECIPIENT NAME] · [ORGANIZATION]\n[ADDRESS]',size=9,after=7,keep=True)
  else:para(d,'[RECIPIENT NAME]\n'+('[LONG RECIPIENT ORGANIZATION NAME FOR PROCUREMENT AND ADMINISTRATIVE CORRESPONDENCE]' if stress else '[RECIPIENT ORGANIZATION]')+'\n[ADDRESS]\n[CITY] · [COUNTRY] · [POSTAL CODE]',after=10,keep=True)
  subj=('Request for supplier information and product documentation' if variant!='Executive' else 'Formal correspondence regarding the proposed review')
  if stress:subj+=' and clarification of available supporting evidence for procurement review and subsequent internal assessment'
  if rtl:para(d,'الموضوع: طلب معلومات ووثائق من المورد'+(' ومراجعة المعلومات المتاحة وتحديد الوثائق الداعمة للمشتريات والتقييم الداخلي' if stress else ''),True,size=12,bold=True,style='Heading 1',after=6,keep=True)
  if lang!='ar':para(d,subj,font='Jost',size=12,bold=True,style='Heading 1',after=12,keep=True)
 if lang=='en':
  
  if not cont:para(d,'Dear [RECIPIENT NAME],',after=10)
  para(d,('We request your review of the accompanying information before a formal response is prepared. Please identify any points that need clarification and the appropriate person to coordinate the review.' if variant=='Executive' else 'Please share the available product information and supporting documentation for our internal review. Indicate which details are confirmed and which require further validation.'))
  para(d,('This letter is sample administrative correspondence. It records no decision, obligation, approval or contractual commitment.' if variant=='Executive' else 'We would welcome a concise response identifying the relevant contact and any information needed from us. This sample does not create a commercial commitment or confirm a supplier relationship.'))
 else:
  
  if not cont:para(d,'السادة «اسم الجهة»،',True,after=6)
  para(d,'يرجى تزويدنا بالمعلومات والوثائق المتاحة للمراجعة الداخلية، مع توضيح ما تم تأكيده وما يحتاج إلى تحقق إضافي.',True)
  if lang=='bi':
   
   if not cont:para(d,'Dear [RECIPIENT NAME],',after=6)
   para(d,'Please share the available information and documentation for our internal review, identifying confirmed details and items requiring further validation.')
  else:para(d,'هذا النص نموذج للمراسلات فقط، ولا يثبت علاقة تجارية أو التزاماً تعاقدياً.',True)
 for page in range(2,pages+1):
  d.add_page_break()
  for i in range(4):para(d,('Additional sample review text. Please identify verified information and any documents requiring further review. This paragraph tests correspondence pagination and remains sample content.' if lang=='en' else 'نص تجريبي إضافي لاختبار الصفحات. يرجى تحديد المعلومات المؤكدة والوثائق التي تحتاج إلى مراجعة إضافية.'),rtl)
 if lang=='en':
  para(d,'Yours sincerely,',before=10,after=4,keep=True)
  p=para(d,'[SIGNATURE]',size=8,after=4,keep=True);p.paragraph_format.space_before=Mm(16)
  para(d,'[SIGNATORY NAME]\n[TITLE]',bold=True,after=0)
 else:
  para(d,'وتفضلوا بقبول التحية،',True,before=5,after=2,keep=True)
  if lang=='bi':para(d,'Yours sincerely,',size=10.5,after=2,keep=True)
  p=para(d,'«التوقيع»',True,size=9,after=2,keep=True);p.paragraph_format.space_before=Mm(12)
  para(d,'«اسم الموقّع»\n«المسمى الوظيفي»',True,bold=True,after=0)
  if lang=='bi':para(d,'[SIGNATORY NAME] · [TITLE]',size=9,after=0)
 if return_document:return d
 path=(R/'04_qa'/'fixtures' if stress or pages>1 else R/'01_templates')/(name+'.docx');path.parent.mkdir(parents=True,exist_ok=True);d.save(path);embed_fonts(path);return path
if __name__=='__main__':
 for lang,display in [('en','English'),('ar','Arabic'),('bi','Bilingual')]:
  for cont in [False,True]:create('GEM_Letterhead_'+display+('_Continuation' if cont else '_First_Page'),lang,cont)
 for v in ['Executive','Minimal']:create('GEM_Letterhead_'+v,variant=v)
 for lang in ['en','ar','bi']:
  for pages in [2,3]:create('QA_'+lang+'_'+str(pages)+'_pages',lang,pages=pages)
  create('QA_'+lang+'_long_fields',lang,stress=True)
