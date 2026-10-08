from build_letterheads import *
from docx.oxml import OxmlElement
for lang,count in [('en',8),('ar',10),('bi',10)]:
 for n in ([8,16,20] if lang=='en' else [count,count+12]):
  display={'en':'English','ar':'Arabic','bi':'Bilingual'}[lang]
  d=Document(R/'01_templates'/('GEM_Letterhead_'+display+'_First_Page.docx'))
  closing=next(p for p in d.paragraphs if p.text in ['Yours sincerely,','وتفضلوا بقبول التحية،'])
  for i in range(n):
   rtl=lang=='ar' or (lang=='bi' and i%2==0)
   p=para(d,('نص تجريبي لاختبار تدفق الصفحات تلقائياً. يرجى مراجعة المعلومات وتحديد ما تم تأكيده وما يحتاج إلى تحقق إضافي قبل إصدار المراسلات.' if rtl else 'SAMPLE CONTENT. This paragraph tests automatic text flow across pages. Please review the available information, identify confirmed details and note any evidence that remains outstanding before correspondence is issued.'),rtl)
   closing._p.addprevious(p._p)
  path=R/'04_qa/fixtures'/('QA_'+lang+'_automatic_'+str(n)+'.docx');d.save(path)
# LTR identifiers and contextual numerals without actual contact/legal claims
for lang in ['ar','bi']:
 display={'ar':'Arabic','bi':'Bilingual'}[lang];d=Document(R/'01_templates'/('GEM_Letterhead_'+display+'_First_Page.docx'))
 p=para(d,'',True);run(p,'اختبار اتجاه المعرفات: ',font='Noto Sans Arabic',arabic=True,size=11.5);run(p,'[EMAIL] [URL] [PHONE] GEM-TEST-2026-001',size=10.5)
 d.paragraphs[1]._p.addnext(p._p)
 p=para(d,'اختبار أرقام سياقية: ١٢٣٤٥٦٧٨٩٠',True);d.paragraphs[2]._p.addnext(p._p)
 path=R/'04_qa/fixtures'/('QA_'+lang+'_bidi.docx');d.save(path)
